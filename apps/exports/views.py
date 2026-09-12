"""Exportación de evidencia a Excel para las auditorías de SUNAFIL.

Cada endpoint devuelve un .xlsx listo para imprimir o adjuntar al expediente. La idea es que
ante una inspección la empresa no tenga que armar nada a mano: el dato ya está en el sistema.
"""
from datetime import timedelta
from io import BytesIO

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from rest_framework.exceptions import PermissionDenied
from rest_framework.views import APIView

from apps.committee.models import Committee
from apps.epp.models import EppDelivery
from apps.inspections.models import Inspection
from apps.iperc.models import IpercEntry
from apps.reports.models import Report

CABECERA = Font(bold=True, color="FFFFFF")
FONDO = PatternFill("solid", fgColor="E65100")


def _hoja(wb: Workbook, titulo: str, columnas: list[str]):
    ws = wb.active if wb.active.max_row == 1 and wb.active.max_column == 1 else wb.create_sheet()
    ws.title = titulo[:31]
    ws.append(columnas)
    for celda in ws[1]:
        celda.font = CABECERA
        celda.fill = FONDO
        celda.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"
    return ws


def _ajustar(ws) -> None:
    for i, columna in enumerate(ws.columns, start=1):
        ancho = max((len(str(c.value)) if c.value is not None else 0) for c in columna)
        ws.column_dimensions[get_column_letter(i)].width = min(max(ancho + 2, 10), 60)


def _responder(wb: Workbook, nombre: str) -> HttpResponse:
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    fecha = timezone.localdate().isoformat()
    response = HttpResponse(
        buffer.read(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{nombre}_{fecha}.xlsx"'
    return response


class _BaseExport(APIView):
    """Solo el supervisor, el comité o un admin pueden exportar la evidencia."""

    def initial(self, request, *args, **kwargs):
        super().initial(request, *args, **kwargs)
        if not request.user.can_manage:
            raise PermissionDenied("Solo supervisor o comité de SST puede exportar evidencia.")

    def _dias(self, request, default: int = 365) -> int:
        try:
            return int(request.query_params.get("days", default))
        except ValueError:
            return default


class ReportsExportView(_BaseExport):
    def get(self, request):
        desde = timezone.now() - timedelta(days=self._dias(request))
        qs = (
            Report.objects.filter(company=request.user.company_id, created_at__gte=desde)
            .select_related("reported_by", "assigned_to", "area", "category")
            .order_by("created_at")
        )
        estado = request.query_params.get("status")
        if estado:
            qs = qs.filter(status=estado)

        wb = Workbook()
        ws = _hoja(
            wb,
            "Actos y condiciones",
            [
                "N°", "Tipo", "Categoría", "Área", "Descripción", "Severidad", "Estado",
                "Reportado por", "Fecha del reporte", "Ocurrió el", "Responsable asignado",
                "Fecha de cierre", "Horas hasta el cierre", "Acción correctiva",
                "Latitud", "Longitud", "Llegó offline",
            ],
        )
        for r in qs:
            ws.append([
                r.id,
                r.get_kind_display(),
                r.category.name if r.category else "",
                r.area.name if r.area else "",
                r.description,
                r.get_severity_display(),
                r.get_status_display(),
                r.reported_by.get_full_name() or r.reported_by.username,
                timezone.localtime(r.created_at).strftime("%Y-%m-%d %H:%M"),
                timezone.localtime(r.occurred_at).strftime("%Y-%m-%d %H:%M"),
                r.assigned_to.get_full_name() if r.assigned_to else "",
                timezone.localtime(r.closed_at).strftime("%Y-%m-%d %H:%M") if r.closed_at else "",
                round(r.resolution_hours, 1) if r.resolution_hours is not None else "",
                r.closure_note,
                float(r.latitude) if r.latitude else "",
                float(r.longitude) if r.longitude else "",
                "Sí" if r.synced_offline else "No",
            ])
        _ajustar(ws)
        return _responder(wb, "registro_actos_condiciones_inseguras")


class IpercExportView(_BaseExport):
    def get(self, request):
        qs = (
            IpercEntry.objects.filter(matrix__company=request.user.company_id)
            .select_related("area", "matrix", "responsible")
            .order_by("area__name", "job_position")
        )
        matriz = request.query_params.get("matrix")
        if matriz:
            qs = qs.filter(matrix_id=matriz)

        wb = Workbook()
        ws = _hoja(
            wb,
            "Matriz IPERC",
            [
                "Versión", "Área", "Puesto / tarea", "Peligro", "Riesgo",
                "Probabilidad", "Consecuencia", "Nivel de riesgo", "Puntaje",
                "Controles existentes", "Controles propuestos", "Responsable",
                "Origen", "Última actualización",
            ],
        )
        for e in qs:
            ws.append([
                f"v{e.matrix.version}",
                e.area.name,
                e.job_position,
                e.hazard,
                e.risk,
                e.get_probability_display(),
                e.get_consequence_display(),
                e.risk_level,
                e.risk_score,
                e.existing_controls,
                e.proposed_controls,
                e.responsible.get_full_name() if e.responsible else "",
                f"Reporte #{e.source_report_id}" if e.source_report_id else "Revisión manual",
                timezone.localtime(e.updated_at).strftime("%Y-%m-%d"),
            ])
        _ajustar(ws)
        return _responder(wb, "matriz_iperc")


class EppExportView(_BaseExport):
    def get(self, request):
        qs = (
            EppDelivery.objects.filter(item__company=request.user.company_id)
            .select_related("item", "worker", "delivered_by")
            .order_by("-delivered_at")
        )
        wb = Workbook()
        ws = _hoja(
            wb,
            "Entrega de EPP",
            [
                "Trabajador", "DNI", "EPP", "Cantidad", "Entregado el", "Entregado por",
                "Vence el", "Vencido", "Conformidad del trabajador", "Observaciones",
            ],
        )
        for d in qs:
            ws.append([
                d.worker.get_full_name() or d.worker.username,
                d.worker.dni,
                d.item.name,
                d.quantity,
                timezone.localtime(d.delivered_at).strftime("%Y-%m-%d"),
                d.delivered_by.get_full_name() if d.delivered_by else "",
                d.expires_at.isoformat() if d.expires_at else "",
                "Sí" if d.is_expired else "No",
                "Sí" if d.acknowledged else "No",
                d.notes,
            ])
        _ajustar(ws)
        return _responder(wb, "registro_entrega_epp")


class InspectionsExportView(_BaseExport):
    def get(self, request):
        qs = (
            Inspection.objects.filter(schedule__company=request.user.company_id)
            .select_related("schedule__area", "performed_by")
            .order_by("-due_date")
        )
        wb = Workbook()
        ws = _hoja(
            wb,
            "Inspecciones",
            [
                "Inspección", "Área", "Frecuencia", "Fecha programada", "Estado",
                "Realizada el", "Realizada por", "Fuera de plazo", "Hallazgos",
            ],
        )
        for i in qs:
            ws.append([
                i.schedule.title,
                i.schedule.area.name,
                i.schedule.get_frequency_display(),
                i.due_date.isoformat(),
                i.get_status_display(),
                timezone.localtime(i.performed_at).strftime("%Y-%m-%d") if i.performed_at else "",
                i.performed_by.get_full_name() if i.performed_by else "",
                "Sí" if i.is_overdue else "No",
                i.findings,
            ])
        _ajustar(ws)
        return _responder(wb, "programa_inspecciones")


class CommitteeExportView(_BaseExport):
    def get(self, request):
        committee = Committee.objects.filter(company=request.user.company_id).first()
        wb = Workbook()
        ws = _hoja(
            wb,
            "Actas del comite",
            [
                "N° de acta", "Fecha", "Tipo", "Lugar", "Asistentes", "Quórum",
                "Agenda", "Acuerdos", "Acuerdos cumplidos",
            ],
        )
        if committee is not None:
            for m in committee.meetings.prefetch_related("agreements", "attendees"):
                acuerdos = list(m.agreements.all())
                ws.append([
                    m.number,
                    m.date.isoformat(),
                    "Extraordinaria" if m.is_extraordinary else "Ordinaria",
                    m.place,
                    m.attendee_count,
                    "Sí" if m.quorum_reached else "No",
                    m.agenda,
                    len(acuerdos),
                    sum(1 for a in acuerdos if a.status == "CUMPLIDO"),
                ])
        _ajustar(ws)
        return _responder(wb, "actas_comite_sst")
