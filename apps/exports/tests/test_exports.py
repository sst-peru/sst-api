from io import BytesIO

import pytest
from django.contrib.auth import get_user_model
from django.urls import reverse
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.accounts.models import Area, Company, Role
from apps.reports.models import Report

User = get_user_model()

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture
def company(db):
    return Company.objects.create(name="Demo SAC", ruc="20100000003", worker_count=45)


@pytest.fixture
def area(company):
    return Area.objects.create(company=company, name="Obra Civil")


@pytest.fixture
def supervisor(company):
    return User.objects.create_user(
        username="sup", password="x", company=company, role=Role.SUPERVISOR
    )


@pytest.fixture
def operario(company, area):
    return User.objects.create_user(
        username="op", password="x", company=company, area=area, role=Role.OPERARIO
    )


def auth(user):
    client = APIClient()
    client.force_authenticate(user=user)
    return client


def test_supervisor_exporta_reportes(supervisor, operario, area, company):
    Report.objects.create(
        company=company, reported_by=operario, kind="CONDICION", area=area,
        description="Cable expuesto en el tablero", severity="ALTA",
    )
    response = auth(supervisor).get(reverse("export-reports"))

    assert response.status_code == 200
    assert response["Content-Type"] == XLSX
    assert "attachment" in response["Content-Disposition"]
    assert ".xlsx" in response["Content-Disposition"]

    # El archivo tiene que abrirse de verdad, no solo tener el content-type correcto.
    wb = load_workbook(BytesIO(response.content))
    ws = wb.active
    assert ws["A1"].value == "N°"
    assert ws.max_row == 2
    assert "Cable expuesto" in str(ws["E2"].value)


def test_operario_no_exporta_evidencia(operario):
    response = auth(operario).get(reverse("export-reports"))
    assert response.status_code == 403


def test_exportaciones_vacias_no_fallan(supervisor):
    """Una empresa sin datos igual debe poder descargar el formato, solo con la cabecera."""
    for nombre in ["export-iperc", "export-epp", "export-inspections", "export-committee"]:
        response = auth(supervisor).get(reverse(nombre))
        assert response.status_code == 200, nombre
        wb = load_workbook(BytesIO(response.content))
        assert wb.active.max_row == 1, nombre


def test_filtro_por_estado_en_la_exportacion(supervisor, operario, area, company):
    Report.objects.create(
        company=company, reported_by=operario, kind="ACTO", area=area, status="ABIERTO"
    )
    cerrado = Report.objects.create(
        company=company, reported_by=operario, kind="ACTO", area=area
    )
    cerrado.close(note="Corregido")

    response = auth(supervisor).get(reverse("export-reports"), {"status": "CERRADO"})
    wb = load_workbook(BytesIO(response.content))
    assert wb.active.max_row == 2  # cabecera + solo el cerrado
